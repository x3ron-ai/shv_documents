from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_BREAK
from docx.shared import Pt, RGBColor, Inches
import os
from docx.oxml import OxmlElement
from docx.enum.style import WD_STYLE_TYPE

class WordDocCreator:
	def __init__(self, filename="document.docx", base_doc=None):
		if base_doc:
			self.document = base_doc
		else:
			self.document = Document()
		self.filename = filename
		self.caption_count = 0

	def add_paragraph(self, text, alignment="left", font_size=12, bold=False, italic=False, 
					 font_name="Arial", color=(0, 0, 0), indent_left=0, indent_first_line=0, line_spacing=1.5):
		paragraph = self.document.add_paragraph()
		run = paragraph.add_run(text)
		alignments = {
			"left": WD_PARAGRAPH_ALIGNMENT.LEFT,
			"center": WD_PARAGRAPH_ALIGNMENT.CENTER,
			"right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
			"justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY
		}
		paragraph.alignment = alignments.get(alignment.lower(), WD_PARAGRAPH_ALIGNMENT.LEFT)
		run.font.size = Pt(font_size)
		run.font.bold = bold
		run.font.italic = italic
		run.font.name = font_name
		run.font.color.rgb = RGBColor(*color)
		# Применяем отступы
		paragraph.paragraph_format.left_indent = Inches(float(indent_left))
		paragraph.paragraph_format.first_line_indent = Inches(float(indent_first_line))
		paragraph.paragraph_format.line_spacing = float(line_spacing)
		# Отладочный вывод
		print(f"Paragraph: Left indent: {indent_left}, First line: {indent_first_line}, Line spacing: {line_spacing}")
		return paragraph

	def add_numbered_list(self, items, start=1, indent_left=0, indent_first_line=0, line_spacing=1.5):
		for i, item in enumerate(items, start):
			paragraph = self.document.add_paragraph(f"{i}. {item}")
			paragraph.style = 'List Number'
			paragraph.paragraph_format.left_indent = Inches(float(indent_left))
			paragraph.paragraph_format.first_line_indent = Inches(float(indent_first_line))
			paragraph.paragraph_format.line_spacing = float(line_spacing)
			# Отладочный вывод
			print(f"Numbered list item: Left indent: {indent_left}, First line: {indent_first_line}, Line spacing: {line_spacing}")
		return paragraph

	def add_bullet_list(self, items, indent_left=0, indent_first_line=0, line_spacing=1.5):
		for item in items:
			paragraph = self.document.add_paragraph(item)
			paragraph.style = 'List Bullet'
			paragraph.paragraph_format.left_indent = Inches(float(indent_left))
			paragraph.paragraph_format.first_line_indent = Inches(float(indent_first_line))
			paragraph.paragraph_format.line_spacing = float(line_spacing)
			# Отладочный вывод
			print(f"Bullet list item: Left indent: {indent_left}, First line: {indent_first_line}, Line spacing: {line_spacing}")
		return paragraph

	def add_table(self, data, col_widths=None):
		rows = len(data)
		cols = len(data[0]) if data else 0
		if rows == 0 or cols == 0:
			return None
		table = self.document.add_table(rows=rows, cols=cols)
		table.style = 'Table Grid'
		if col_widths and len(col_widths) == cols:
			for i, width in enumerate(col_widths):
				for cell in table.columns[i].cells:
					cell.width = Inches(width)
		for i, row in enumerate(data):
			for j, value in enumerate(row):
				table.rows[i].cells[j].text = str(value)
		return table

	def add_picture(self, image_path, width=None, height=None, align="left"):
		if not os.path.exists(image_path):
			raise FileNotFoundError(f"Файл изображения {image_path} не найден")
		paragraph = self.document.add_paragraph()
		paragraph.alignment = {
			"left": WD_PARAGRAPH_ALIGNMENT.LEFT,
			"center": WD_PARAGRAPH_ALIGNMENT.CENTER,
			"right": WD_PARAGRAPH_ALIGNMENT.RIGHT
		}.get(align.lower(), WD_PARAGRAPH_ALIGNMENT.LEFT)
		run = paragraph.add_run()
		if width and height:
			run.add_picture(image_path, width=Inches(width), height=Inches(height))
		elif width:
			run.add_picture(image_path, width=Inches(width))
		elif height:
			run.add_picture(image_path, height=Inches(height))
		else:
			run.add_picture(image_path)
		paragraph.paragraph_format.left_indent = Inches(0)
		paragraph.paragraph_format.right_indent = Inches(0)
		return paragraph

	def add_caption(self, caption_text, label="Рисунок", bold=False, size=12, face="Times New Roman", color=(0, 0, 0)):
		"""Добавляет подпись к последнему изображению с указанными стилями."""
		self.caption_count += 1
		caption = f"{label} {self.caption_count} - {caption_text}"
		paragraph = self.document.add_paragraph()
		run = paragraph.add_run(caption)
		run.font.bold = bold
		run.font.size = Pt(size)
		run.font.name = face
		run.font.color.rgb = RGBColor(*color)
		paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
		paragraph.style = 'Caption'
		return paragraph

	def save(self):
		try:
			self.document.save(self.filename)
			return f"Документ сохранен как {os.path.abspath(self.filename)}"
		except Exception as e:
			return f"Ошибка при сохранении: {str(e)}"
