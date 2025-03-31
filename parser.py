from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import xml.etree.ElementTree as ET
from docx.shared import Pt, Inches, RGBColor

class XMLToWordParser:
    def __init__(self, xml_path, output_path, title_page_path=None):
        self.xml_path = xml_path
        self.output_path = output_path
        self.title_page_path = title_page_path
        self.doc = Document()

    def parse_and_convert(self):
        try:
            if self.title_page_path:
                title_doc = Document(self.title_page_path)
                for element in title_doc.element.body:
                    self.doc.element.body.append(element)

            with open(self.xml_path, 'r', encoding='utf-8') as f:
                root = ET.fromstring(f.read())

            for element in root:
                if element.tag == "text":
                    self._add_text(element)
                elif element.tag == "table":
                    self._add_table(element)
                elif element.tag == "list":
                    self._add_list(element)
                elif element.tag == "image":
                    self._add_image(element)

            self.doc.save(self.output_path)
            return self.output_path
        except Exception as e:
            return f"Ошибка при генерации документа: {str(e)}"

    def _add_text(self, text_element):
        text = text_element.text.strip() if text_element.text else ""
        alignment = text_element.get("align", "justify")
        indent_left_cm = float(text_element.get("indent_left", "0"))
        indent_first_line_cm = float(text_element.get("indent_first_line", "0"))
        line_spacing = float(text_element.get("line_spacing", "1.5"))
        font_face = text_element.get("font_face", "Times New Roman")
        font_size = int(text_element.get("font_size", "14"))

        align_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
        }

        if text:
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.alignment = align_map.get(alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            paragraph.paragraph_format.left_indent = Inches(indent_left_cm * 0.393701)
            paragraph.paragraph_format.first_line_indent = Inches(indent_first_line_cm * 0.393701)
            paragraph.paragraph_format.line_spacing = line_spacing
            parts = re.split(r'(<[^>]+>)', text)
            bold = False
            font_color = (0, 0, 0)

            for part in parts:
                if part.startswith('<') and part.endswith('>'):
                    if part == '<b>':
                        bold = True
                    elif part == '</b>':
                        bold = False
                    elif part.startswith('<font'):
                        attrs = re.findall(r'(\w+)="([^"]+)"', part)
                        for attr, value in attrs:
                            if attr == "size":
                                font_size = int(value)
                            elif attr == "color":
                                font_color = tuple(int(value[i:i+2], 16) for i in (0, 2, 4))
                    elif part == '</font>':
                        font_color = (0, 0, 0)
                else:
                    if part.strip():
                        run = paragraph.add_run(part)
                        run.font.bold = bold
                        run.font.size = Pt(font_size)
                        run.font.name = font_face
                        run.font.color.rgb = RGBColor(*font_color)

    def _add_table(self, table_element):
        col_widths = [float(x.strip()) for x in table_element.get("col_widths", "2,1,2").split(',')]
        rows = list(table_element.findall("row"))
        table = self.doc.add_table(rows=len(rows), cols=len(col_widths))
        table.style = 'Table Grid'

        for i, row in enumerate(rows):
            cells = row.findall("cell")
            for j, cell in enumerate(cells):
                table.rows[i].cells[j].text = cell.text or ""

        for i, width in enumerate(col_widths):
            for cell in table.columns[i].cells:
                cell.width = Inches(width)

    def _add_list(self, list_element):
        list_type = list_element.get("type", "bullet")
        items = list_element.findall("item")
        for item in items:
            if item.text:
                p = self.doc.add_paragraph(item.text)
                p.style = 'List Bullet' if list_type == "bullet" else 'List Number'

    def _add_image(self, image_element):
        path = image_element.get("path")
        if path and os.path.exists(path):
            self.doc.add_picture(path, width=Inches(4))
            caption = image_element.get("caption")
            if caption:
                p = self.doc.add_paragraph(caption)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.runs[0]
                run.font.bold = image_element.get("caption_bold") == "true"
                run.font.size = Pt(int(image_element.get("caption_size", "12")))
                run.font.name = image_element.get("caption_face", "Times New Roman")
                color = image_element.get("caption_color", "000000")
                run.font.color.rgb = RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
